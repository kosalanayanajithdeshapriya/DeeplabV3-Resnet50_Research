 """Convert paper_report/PAPER_RESULTS_REPORT.md into an editable .docx with embedded figures.
Handles the Markdown subset the report uses: headings, paragraphs, **bold**/*italic*/`code`,
bullet and numbered lists, block quotes, fenced code, pipe tables, and images."""
import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor
from PIL import Image

ROOT = Path(__file__).resolve().parents[1] / "paper_report"
# optional: python build_paper_report_docx.py OTHER_REPORT.md  (resolved inside paper_report/)
SRC = ROOT / (sys.argv[1] if len(sys.argv) > 1 else "PAPER_RESULTS_REPORT.md")
OUT = SRC.with_suffix(".docx")

doc = Document()
sec = doc.sections[0]
sec.left_margin = sec.right_margin = Cm(2.0)
sec.top_margin = sec.bottom_margin = Cm(2.0)
normal = doc.styles["Normal"]
normal.font.name = "Calibri"
normal.font.size = Pt(10.5)
TEXT_W = sec.page_width - sec.left_margin - sec.right_margin

INLINE = re.compile(r"(\*\*[^*]+\*\*|\*[^*\s][^*]*\*|`[^`]+`)")


def add_runs(par, text, size=None):
    for tok in INLINE.split(text):
        if not tok:
            continue
        if tok.startswith("**"):
            r = par.add_run(tok[2:-2]); r.bold = True
        elif tok.startswith("`"):
            r = par.add_run(tok[1:-1]); r.font.name = "Consolas"
        elif tok.startswith("*") and tok.endswith("*") and len(tok) > 2:
            r = par.add_run(tok[1:-1]); r.italic = True
        else:
            r = par.add_run(tok)
        if size:
            r.font.size = Pt(size)


def shade(cell, hex_fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), hex_fill)
    tcPr.append(shd)


def add_table(rows):
    cells = [[c.strip() for c in r.strip().strip("|").split("|")] for r in rows]
    header, body = cells[0], [r for r in cells[2:]]
    ncol = len(header)
    t = doc.add_table(rows=1 + len(body), cols=ncol)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    size = 8 if ncol >= 6 else 9
    for j, h in enumerate(header):
        c = t.rows[0].cells[j]; c.text = ""
        add_runs(c.paragraphs[0], f"**{h}**" if h and not h.startswith("**") else h, size)
        shade(c, "DCE8F5")
    for i, r in enumerate(body, start=1):
        for j in range(ncol):
            c = t.rows[i].cells[j]; c.text = ""
            add_runs(c.paragraphs[0], r[j] if j < len(r) else "", size)
    doc.add_paragraph()


lines = SRC.read_text(encoding="utf-8").splitlines()
i = 0
para_buf = []


def flush():
    global para_buf
    if para_buf:
        add_runs(doc.add_paragraph(), " ".join(para_buf))
        para_buf = []


while i < len(lines):
    ln = lines[i]
    s = ln.strip()
    if s.startswith("```"):
        flush()
        i += 1
        code = []
        while not lines[i].strip().startswith("```"):
            code.append(lines[i]); i += 1
        p = doc.add_paragraph()
        r = p.add_run("\n".join(code)); r.font.name = "Consolas"; r.font.size = Pt(8)
    elif s.startswith("|"):
        flush()
        block = []
        while i < len(lines) and lines[i].strip().startswith("|"):
            block.append(lines[i]); i += 1
        add_table(block)
        continue
    elif m := re.match(r"!\[[^\]]*\]\(([^)]+)\)", s):
        flush()
        img = ROOT / m.group(1)
        with Image.open(img) as im:  # small figures keep their natural size (150 dpi) instead of full width
            width = min(TEXT_W, Inches(im.width / 150))
        doc.add_picture(str(img), width=width)
        doc.paragraphs[-1].alignment = 1
    elif m := re.match(r"(#{1,4})\s+(.*)", s):
        flush()
        level = len(m.group(1))
        doc.add_heading(m.group(2).replace("**", ""), level=0 if level == 1 else level - 1)
    elif s == "---":
        flush()
    elif s.startswith(">"):
        flush()
        p = doc.add_paragraph(style="Intense Quote")
        add_runs(p, s.lstrip("> ").strip())
    elif m := re.match(r"(\s*)- (.*)", ln):
        flush()
        style = "List Bullet 2" if len(m.group(1)) >= 2 else "List Bullet"
        add_runs(doc.add_paragraph(style=style), m.group(2))
    elif m := re.match(r"\d+\.\s+(.*)", s):
        flush()
        add_runs(doc.add_paragraph(style="List Number"), m.group(1))
    elif not s:
        flush()
    elif s.startswith("*Figure"):
        flush()
        p = doc.add_paragraph()
        add_runs(p, s, 9)
        for r in p.runs:
            r.font.color.rgb = RGBColor(0x52, 0x51, 0x4E)
    else:
        para_buf.append(s)
    i += 1
flush()
doc.save(OUT)
print("wrote", OUT)
